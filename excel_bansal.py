# -*- coding: utf-8 -*-
"""Planilha da replicação de Bansal et al. (2022).

Contas, deflator, agências, painel, direções, restrições, índices e tabelas são fórmula.
β, λ e φ saem de solucoes.xlsx, gravado por `python excel_bansal.py --solucionar`.
O paper resolve no LINGO e não descreve o algoritmo. Não há macro nem Python dentro do Excel.
"""
import os
import re
import sqlite3
import sys
import zipfile
from xml.etree import ElementTree as ET

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

import dsmlpi as m
import painel_bansal as pb

NEG = Font(bold=True)
CAB = PatternFill("solid", fgColor="DDEBF7")
def _cab(ws, linha, textos, col=1):
    for i, t in enumerate(textos):
        c = ws.cell(linha, col + i, t)
        c.font = NEG
        c.fill = CAB
        c.alignment = Alignment(wrap_text=False, vertical="top")


def _val(x):
    if x is None:
        return None
    try:
        return None if not np.isfinite(x) else float(x)
    except TypeError:
        return x


ARQUIVO_SOLUCOES = "solucoes.xlsx"


def _leiame(ws):
    ws.title = "Leia-me"
    textos = [
        "Replicação de Bansal, Kumar, Mehra e Gulati (2022), Omega 107:102538.",
        "Contas: saldos do IF.data (conglomerado financeiro). Fluxos da DRE somam junho e dezembro, porque a DRE é acumulada no semestre. Se o código do conglomerado começa no meio da janela, os semestres anteriores vêm do CNPJ do banco comercial, quando o ativo total varia no máximo 15% na troca.",
        "Descricao: nome e descrição de cada conta no dicionário do IF.data. A coluna da descrição é a fórmula do COSIF (Plano Contábil das Instituições do Sistema Financeiro Nacional). O nome de cada conta COSIF está na coluna ao lado e na tabela abaixo.",
        "Deflator: SGS 1211 (deflator implícito do PIB), índice com 2024 = 1, montado por fórmula a partir da variação anual.",
        "Agencias: soma de AGEN_PROCESSADAS da ESTBAN por conglomerado financeiro (dezembro).",
        f"Painel: variáveis da Tabela 1 por fórmula a partir de Contas; normalizadas = nominal / índice / {pb.UNIDADE:,.0f} / agências (trabalho sem dividir por agências).",
        "Fronteiras e Direcoes: fórmulas. A direção usa MAX e MIN com as constantes do paper (ξ, π, η, δ).",
        "Solucao: ligação para solucoes.xlsx, na mesma pasta. β, o status e os pesos (λ, φ ou z) são número nesse arquivo. status 0 é ótimo.",
        "O comando python excel_bansal.py --solucionar grava solucoes.xlsx. Ele roda o simplex em duas fases, em Python, sem biblioteca. O paper resolve no LINGO e não descreve o algoritmo. Não há macro nem Python dentro do Excel.",
        "Restricoes: lado esquerdo, lado direito e folga por fórmula. Folga ≥ 0 (ou 0 na igualdade) prova a restrição.",
        "Indices: para cada banco e par de anos, as quatro distâncias, a mudança de eficiência (EC), a mudança técnica (TC) e o índice.",
        "Tabelas: média geométrica dos quatro índices, dos bancos juntos, por ano e por banco. O MLPI (Malmquist-Luenberger, sem rede) e o SMLPI (sequencial, sem rede) usam só quem tem margem de juros positiva em todos os anos.",
        "Ao abrir esta pasta, atualize as ligações. Sem solucoes.xlsx na mesma pasta, a aba Solucao fica sem valor.",
    ]
    for i, t in enumerate(textos, 1):
        ws.cell(i, 1, t)
    ws.column_dimensions["A"].width = 140


def _contas(ws, contas_long, ordem):
    por = {(r.sigla, r.ano, r.mes, r.conta): r.saldo for r in contas_long.itertuples()}
    cab = ["sigla", "ano"]
    colunas = []
    for conta, (rel, tipo) in pb.CONTAS.items():
        if tipo == "fluxo":
            colunas += [(conta, "06"), (conta, "12"), (conta, "anual")]
            cab += [f"{conta} jun (rel {rel})", f"{conta} dez (rel {rel})", f"{conta} anual"]
        else:
            colunas += [(conta, "12")]
            cab += [f"{conta} dez (rel {rel})"]
    _cab(ws, 1, cab)
    col_anual = {}
    for i, (conta, mes) in enumerate(colunas):
        c = i + 3
        if mes == "anual" or pb.CONTAS[conta][1] == "saldo":
            col_anual[conta] = L(c)
    for r, (ano, cod) in enumerate(ordem, 2):
        ws.cell(r, 1, cod); ws.cell(r, 2, int(ano))
        for i, (conta, mes) in enumerate(colunas):
            c = i + 3
            if mes == "anual":
                ws.cell(r, c, f"={L(c - 2)}{r}+{L(c - 1)}{r}")
            else:
                ws.cell(r, c, _val(por.get((cod, ano, mes, conta))))
    return col_anual


# Nomes do elenco COSIF (Plano Contábil das Instituições do Sistema Financeiro Nacional)
# para as contas que o IF.data cita entre colchetes na descrição.
NOMES_COSIF = {
    "10000007": "Ativo Realizável",
    "13000004": "Títulos e Valores Mobiliários e Instrumentos Financeiros Derivativos",
    "16000001": "Operações de Crédito",
    "17000000": "Operações de Arrendamento Mercantil",
    "20000004": "Ativo Permanente",
    "23000001": "Imobilizado de Arrendamento",
    "41000007": "Depósitos",
    "41300006": "Depósitos Interfinanceiros",
    "46000002": "Obrigações por Empréstimos e Repasses",
    "49908008": "Credores por Antecipação de Valor Residual",
    "60000002": "Patrimônio Líquido",
    "70000009": "Resultado Credor",
    "71700009": "Rendas de Prestação de Serviços",
    "71794008": "Rendas de Pacotes de Serviços – PF",
    "71795007": "Rendas de Serviços Prioritários – PF",
    "71796006": "Rendas de Serviços Diferenciados – PF",
    "71797005": "Rendas de Serviços Especiais – PF",
    "71798004": "Rendas de Tarifas Bancárias – PJ",
    "71930006": "Recuperação de Encargos e Despesas",
    "71970004": "Rendas de Garantias Prestadas",
    "71990307": "Operações de Crédito de Liquidação Duvidosa",
    "71990352": "Repasses Interfinanceiros",
    "71990400": "Créditos de Arrendamento de Liquidação Duvidosa",
    "71990503": "Perdas na Venda de Valor Residual",
    "71990606": "Outros Créditos de Liquidação Duvidosa",
    "80000006": "Resultado Devedor",
    "8170006": "Despesas Administrativas",
    "81718005": "Despesas de Honorários",
    "81727003": "Despesas de Pessoal – Benefícios",
    "81730007": "Despesas de Pessoal – Encargos Sociais",
    "81733004": "Despesas de Pessoal – Proventos",
    "81736001": "Despesas de Pessoal – Treinamento",
    "81737000": "Despesas de Remuneração de Estagiários",
    "81800009": "Aprovisionamentos e Ajustes Patrimoniais",
    "81810006": "Despesas de Amortização",
    "81820003": "Despesas de Depreciação",
    "81830055": "Perdas em Aplicações em Depósitos Interfinanceiros",
    "81830103": "Desvalorização de Títulos Livres",
    "81830127": "Desvalorização de Créditos Vinculados",
    "81830158": "Desvalorização de Títulos Vinculados a Operações Compromissadas",
    "81830206": "Desvalorização de Títulos Vinculados a Negociação e Intermediação de Valores",
    "81830268": "Derivativos de Crédito",
    "81830309": "Provisões para Operações de Crédito",
    "81830354": "Repasses Interfinanceiros",
    "81830402": "Provisões para Arrendamento Mercantil",
    "81830505": "Perdas na Venda de Valor Residual",
    "81830550": "Perdas de Bens de Arrendamento Operacional",
    "81830608": "Provisões para Outros Créditos",
    "81830701": "Perdas em Participações Societárias",
    "81830804": "Perdas em Dependências no Exterior",
    "81830907": "Perdas em Sociedades Coligadas e Controladas",
    "81900002": "Outras Despesas Operacionais",
    "81910009": "Despesas de Administração de Fundos e Programas Sociais",
    "81912007": "Despesas de Obrigações por Operações Vinculadas à Cessão",
    "81915004": "Prejuízos em Operações de Venda ou de Transferência de Ativos Financeiros",
    "81925001": "Despesas de Imposto sobre Serviços de Qualquer Natureza – ISS",
    "81930003": "Despesas de Contribuição ao COFINS",
    "81933000": "Despesas de Contribuição ao PIS/PASEP",
    "81940000": "Despesas de Cessão de Créditos de Arrendamento",
    "81945005": "Despesas de Cessão de Créditos Decorrentes de Contratos de Exportação",
    "81950007": "Despesas de Cessão de Operações de Crédito",
    "81960004": "Despesas de Obrigações por Fundos Financeiros e de Desenvolvimento",
    "81986002": "Dispêndios de Depósitos Intercooperativos",
    "81990108": "Impostos e Contribuições Sobre Lucros",
    "81990201": "Impostos e Contribuições Sobre Salários",
    "81990304": "Impostos e Contribuições Sobre Serviços de Terceiros",
    "81990902": "Outros",
}


def _codigo_cosif(codigo):
    d = re.sub(r"\D", "", str(codigo))
    if len(d) == 7:
        d = d[:3] + "0" + d[3:]
    if len(d) != 8:
        return ""
    return f"{d[0]}.{d[1]}.{d[2]}.{d[3:5]}.{d[5:7]}-{d[7]}"


def _nome_cosif(codigo):
    nome = NOMES_COSIF.get(str(codigo), "")
    if nome:
        return nome
    if _codigo_cosif(codigo):
        return "sem título no elenco COSIF vigente"
    return ""


def _refs_cosif(desc):
    return re.findall(r"\[(\d+)\]", desc or "")


def _texto_refs_cosif(desc):
    linhas = []
    for codigo in _refs_cosif(desc):
        pont = _codigo_cosif(codigo)
        nome = _nome_cosif(codigo)
        if pont:
            linhas.append(f"{codigo}  {pont}  {nome}")
        else:
            linhas.append(f"{codigo}  {nome}")
    return "\n".join(linhas)


def _dicionario(ws):
    """Nome e descrição oficiais do IF.data, com o nome COSIF de cada conta da fórmula."""
    _cab(ws, 1, ["conta", "relatório", "nome do relatório", "grupo",
                 "nome no Banco Central do Brasil", "descrição no Banco Central do Brasil",
                 "contas COSIF da descrição", "tipo", "variável do modelo"])
    try:
        con = pb.abrir_fonte()
        dic = {(str(rel), str(conta)): (grupo, nome, desc, nome_rel)
               for rel, nome_rel, conta, grupo, nome, desc in con.execute(
                   "SELECT c.num_relatorio, r.nome, c.conta, c.grupo, c.nome_coluna, c.descricao_coluna "
                   "FROM contas c LEFT JOIN relatorios r ON r.numero = c.num_relatorio")}
        con.close()
    except sqlite3.Error:
        dic = {}
    uso = {}
    for var, (_n, _f, termos, _) in pb.VARIAVEIS.items():
        for _s, conta in termos:
            uso.setdefault(conta, []).append(var)
    citados = {}
    r = 2
    for conta, (rel, tipo) in pb.CONTAS.items():
        grupo, nome, desc, nome_rel = dic.get((rel, conta), (None, None, None, None))
        refs = _texto_refs_cosif(desc)
        ws.cell(r, 1, conta)
        ws.cell(r, 2, rel)
        ws.cell(r, 3, nome_rel)
        ws.cell(r, 4, grupo)
        c_nome = ws.cell(r, 5, nome)
        c_desc = ws.cell(r, 6, desc)
        c_refs = ws.cell(r, 7, refs or None)
        for cel in (c_nome, c_desc, c_refs):
            cel.alignment = Alignment(wrap_text=False, vertical="top")
        ws.cell(r, 8, tipo)
        ws.cell(r, 9, ", ".join(uso.get(conta, [])))
        for codigo in _refs_cosif(desc):
            citados.setdefault(codigo, []).append(conta)
        r += 1
    r += 1
    ws.cell(r, 1, "Contas do COSIF (Plano Contábil das Instituições do Sistema Financeiro Nacional) citadas na descrição").font = NEG
    r += 1
    _cab(ws, r, ["conta COSIF", "código no COSIF", "nome no COSIF", "contas do IF.data"])
    r += 1
    for codigo in sorted(citados, key=lambda x: x.zfill(8)):
        ws.cell(r, 1, codigo)
        ws.cell(r, 2, _codigo_cosif(codigo))
        ws.cell(r, 3, _nome_cosif(codigo))
        ws.cell(r, 4, ", ".join(citados[codigo]))
        r += 1
    ws.column_dimensions["C"].width = 62
    ws.column_dimensions["D"].width = 36
    ws.column_dimensions["E"].width = 36
    ws.column_dimensions["F"].width = 70
    ws.column_dimensions["G"].width = 78
    ws.column_dimensions["I"].width = 22


def _deflator(ws, defl):
    _cab(ws, 1, ["ano", "variação % (SGS 1211)", "índice (2024 = 1)"])
    d = defl.sort_values("ano").reset_index(drop=True)
    n = len(d)
    for i, r in d.iterrows():
        lin = i + 2
        ws.cell(lin, 1, int(r["ano"]))
        ws.cell(lin, 2, _val(r["variacao_pct"]))
        if int(r["ano"]) == pb.ANO_BASE_PRECOS:
            ws.cell(lin, 3, 1)
        else:
            ws.cell(lin, 3, f"=C{lin + 1}/(1+B{lin + 1}/100)")
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 18
    return n


def _agencias(ws, ag):
    _cab(ws, 1, ["sigla", "ano", "agências (ESTBAN)", "agências esperadas", "CNPJs somados"])
    for i, r in enumerate(ag.sort_values(["ano", "sigla"]).itertuples(), 2):
        ws.cell(i, 1, r.sigla); ws.cell(i, 2, int(r.ano)); ws.cell(i, 3, int(r.agencias))
        ws.cell(i, 4, int(r.agen_esperadas)); ws.cell(i, 5, r.cnpjs)


def _painel(ws, norm, ordem, col_conta):
    n_bancos = norm["sigla"].nunique()
    info = norm.set_index(["ano", "sigla"])
    cab = ["sigla", "ano", "grupo", "S", "índice deflator", "agências"]
    vs = pb.VARS_MODELO
    cab += [f"{v} nominal" for v in vs] + [f"{v} normalizado" for v in vs]
    cab += ["npl_lag normalizado", "unused_lag normalizado", "profit_lag normalizado"]
    _cab(ws, 1, cab)
    c_nom = {v: 7 + i for i, v in enumerate(vs)}
    c_norm = {v: 7 + len(vs) + i for i, v in enumerate(vs)}
    c_lag = {"npl_lag": 7 + 2 * len(vs), "unused_lag": 8 + 2 * len(vs), "profit_lag": 9 + 2 * len(vs)}
    for r, (ano, cod) in enumerate(ordem, 2):
        row = info.loc[(ano, cod)]
        ws.cell(r, 1, cod); ws.cell(r, 2, int(ano))
        ws.cell(r, 3, row["grupo"]); ws.cell(r, 4, row["sr"])
        ws.cell(r, 5, f"=INDEX(Deflator!$C:$C,MATCH(B{r},Deflator!$A:$A,0))")
        ws.cell(r, 6, f"=SUMIFS(Agencias!$C:$C,Agencias!$A:$A,A{r},Agencias!$B:$B,B{r})")
        for v in vs:
            _, _, termos, usa_abs = pb.VARIAVEIS[v]
            expr = "".join(f"{'+' if s > 0 else '-'}Contas!{col_conta[c]}{r}" for s, c in termos)
            expr = expr.lstrip("+")
            ws.cell(r, c_nom[v], f"=ABS({expr})" if usa_abs else f"={expr}")
            den = f"/E{r}/{pb.UNIDADE:.0f}" + ("" if v in pb.VARS_SEM_NORMALIZAR else f"/F{r}")
            ws.cell(r, c_norm[v], f"={L(c_nom[v])}{r}{den}")
        if int(ano) > pb.ANOS[0]:
            for lag, base in (("npl_lag", "npl"), ("unused_lag", "unused"), ("profit_lag", "profit")):
                ws.cell(r, c_lag[lag], f"={L(c_norm[base])}{r - n_bancos}")
    col = {v: L(c) for v, c in c_norm.items()}
    col.update({k: L(c) for k, c in c_lag.items()})
    return col


def _faixa(col_ini, n, linha):
    return f"{L(col_ini)}{linha}:{L(col_ini + n - 1)}{linha}"


def _modelo_seq(modelo):
    return modelo in ("DSMLPI", "SMLPI")


def _bancos_do_modelo(prep, modelo):
    if modelo in ("MLPI", "SMLPI"):
        js, _ = m.amostra_tabela5(prep)
        return [prep["bancos"][j] for j in js]
    return list(prep["bancos"])


def _linhas_lp(prep):
    """Ordem única das linhas de Solucao, na planilha e em solucoes.xlsx."""
    anos = prep["anos"]
    linhas = []
    for modelo in ("DMLPI", "DSMLPI", "MLPI", "SMLPI"):
        for it in range(len(anos) - 1):
            t = int(anos[it])
            for cod in _bancos_do_modelo(prep, modelo):
                for a in (0, 1):
                    for b in (0, 1):
                        linhas.append((modelo, cod, t, a, b))
    return linhas


def _n_pesos(kind, n):
    return n if kind == "ml" else 4 * n


def _constantes(modelo, a):
    if modelo == "DMLPI":
        return m.XI, m.PI
    return (m.ETA_T if a == 0 else m.ETA_T1), m.DELTA


def _fronteiras(wb, prep, linha_de, col_norm):
    wf = wb.create_sheet("Fronteiras")
    wd = wb.create_sheet("Direcoes")
    _cab(wf, 1, ["modelo", "t", "a", "k", "ano", "sigla"] + list(m.CAMPOS))
    _cab(wd, 1, ["modelo", "t", "a", "grupo", "variavel", "mult", "sub", "valor"])
    fr = {}
    dr = {}
    rf = 2
    rd = 2
    anos = prep["anos"]
    for modelo in ("DMLPI", "DSMLPI", "MLPI", "SMLPI"):
        bancos = _bancos_do_modelo(prep, modelo)
        seq = _modelo_seq(modelo)
        for it in range(len(anos) - 1):
            t = int(anos[it])
            for a in (0, 1):
                idx = list(range(it + a + 1)) if seq else [it + a]
                pontos = [(int(anos[i]), cod) for i in idx for cod in bancos]
                ini = rf
                for k, (ano, cod) in enumerate(pontos):
                    pr = linha_de[(ano, cod)]
                    wf.cell(rf, 1, modelo)
                    wf.cell(rf, 2, t)
                    wf.cell(rf, 3, a)
                    wf.cell(rf, 4, k)
                    wf.cell(rf, 5, ano)
                    wf.cell(rf, 6, cod)
                    for j, v in enumerate(m.CAMPOS):
                        wf.cell(rf, 7 + j, f"=Painel!{col_norm[v]}{pr}")
                    rf += 1
                fr[(modelo, t, a)] = (ini, rf - 1, len(pontos))
                if modelo in ("MLPI", "SMLPI"):
                    continue
                mult, sub = _constantes(modelo, a)
                dini = rd
                mapa = {}
                for grupo, var in m.ORDEM_DIR:
                    c = 7 + m.CAMPOS.index(var)
                    rng = f"Fronteiras!{L(c)}{ini}:{L(c)}{rf - 1}"
                    wd.cell(rd, 1, modelo)
                    wd.cell(rd, 2, t)
                    wd.cell(rd, 3, a)
                    wd.cell(rd, 4, grupo)
                    wd.cell(rd, 5, var)
                    wd.cell(rd, 6, mult)
                    wd.cell(rd, 7, sub)
                    if grupo in ("mu", "omega", "alpha"):
                        wd.cell(rd, 8, f"=F{rd}*MAX(MAX({rng}),-MIN({rng}))")
                    else:
                        wd.cell(rd, 8, f"=MIN({rng})-G{rd}")
                    mapa[var] = rd
                    rd += 1
                dr[(modelo, t, a)] = (dini, rd - 1, mapa)
    return fr, dr


def _solucao(wb, prep, fr, dr):
    wo = wb.create_sheet("Obs")
    ws = wb.create_sheet("Solucao")
    _cab(wo, 1, list(m.CAMPOS))
    _cab(ws, 1, ["modelo", "sigla", "t", "a", "b", "β", "status", "pesos"])
    lookup = {}
    regs = []
    r = 2
    d_padrao = next(iter(dr.values()), None)
    for modelo, cod, t, a, b in _linhas_lp(prep):
        kind = "ml" if modelo in ("MLPI", "SMLPI") else "rede"
        f1, f2, n = fr[(modelo, t, a)]
        if kind == "rede":
            mapa = dr[(modelo, t, a)][2]
        else:
            mapa = None if d_padrao is None else d_padrao[2]
        ws.cell(r, 1, modelo)
        ws.cell(r, 2, cod)
        ws.cell(r, 3, t)
        ws.cell(r, 4, a)
        ws.cell(r, 5, b)
        for c in range(6, 8 + _n_pesos(kind, n)):
            ws.cell(r, c, f"=[1]Solucao!{L(c)}{r}")
        regs.append({"modelo": modelo, "cod": cod, "t": t, "a": a, "b": b, "row": r,
                     "kind": kind, "f1": f1, "f2": f2, "n": n, "dir": mapa})
        lookup[(modelo, cod, t, a, b)] = r
        r += 1
    return lookup, regs


def _obs(wb, prep, lookup, linha_de, col_norm):
    wo = wb["Obs"]
    anos = prep["anos"]
    for (modelo, cod, t, a, b), r in lookup.items():
        it = anos.index(t)
        ano = int(anos[it + b])
        pr = linha_de[(ano, cod)]
        for j, v in enumerate(m.CAMPOS):
            wo.cell(r, 1 + j, f"=Painel!{col_norm[v]}{pr}")


def _linhas_rede(reg):
    rs, ro, n = reg["row"], reg["row"], reg["n"]
    f1, f2 = reg["f1"], reg["f2"]
    beta = f"Solucao!F{rs}"

    def fr(v):
        c = 7 + m.CAMPOS.index(v)
        return f"Fronteiras!{L(c)}{f1}:{L(c)}{f2}"

    def peso(div):
        ini = 8 + (div - 1) * n
        return f"Solucao!{_faixa(ini, n, rs)}"

    phi = f"Solucao!{_faixa(8 + 3 * n, n, rs)}"

    def sp(w, v):
        return f"SUMPRODUCT({w},{fr(v)})"

    def o(v):
        return f"Obs!{L(m.CAMPOS.index(v) + 1)}{ro}"

    def d(v):
        return f"Direcoes!H{reg['dir'][v]}"

    # M2 repete a forma do M1 nas eq. (17)-(25).
    base = 12 if reg["modelo"] == "DSMLPI" else 0

    def eq(n):
        return f"({n + base})"

    linhas = []
    for div, vs in m.X_DIV.items():
        for v in vs:
            linhas.append((eq(5), v, f"={sp(peso(div), v)}", f"={o(v)}-{beta}*({o(v)}+{d(v)})", "≤"))
    for v in m.V1:
        linhas.append((eq(6), v, f"={sp(peso(1), v)}+SUMPRODUCT({phi},{fr(v)})", f"={o(v)}", "≤"))
    for v in m.W2:
        linhas.append((eq(7), v, f"={sp(peso(2), v)}", f"={o(v)}+{beta}*({o(v)}-{d(v)})", "≥"))
    for v in m.Y3:
        linhas.append((eq(8), v, f"={sp(peso(3), v)}", f"={o(v)}+{beta}*({o(v)}-{d(v)})", "≥"))
    for v in m.U2:
        linhas.append((eq(9), v, f"={sp(peso(2), v)}", f"={o(v)}-{beta}*({o(v)}+{d(v)})", "≤"))
    for v in m.C2:
        linhas.append((eq(10), v, f"={sp(peso(2), v)}", f"={o(v)}+{beta}*({o(v)}-{d(v)})", "≥"))
    for v in m.C2_LAG:
        linhas.append((eq(11), v, f"={sp(peso(2), v)}", f"={o(v)}-{beta}*({o(v)}+{d(v)})", "≤"))
    for v in m.Z12:
        linhas.append((eq(12), v, f"={sp(peso(1), v)}-{sp(peso(2), v)}", "=0", "="))
    for v in m.Z23:
        linhas.append((eq(12), v, f"={sp(peso(2), v)}-{sp(peso(3), v)}", "=0", "="))
    for div in (1, 2, 3):
        linhas.append((eq(13), f"Σλ{div}", f"=SUM({peso(div)})", "=1", "="))
    return linhas


def _linhas_ml(reg):
    rs, n = reg["row"], reg["n"]
    f1, f2 = reg["f1"], reg["f2"]
    beta = f"Solucao!F{rs}"
    z = f"Solucao!{_faixa(8, n, rs)}"

    def fr(v):
        c = 7 + m.CAMPOS.index(v)
        return f"Fronteiras!{L(c)}{f1}:{L(c)}{f2}"

    def sp(v):
        return f"SUMPRODUCT({z},{fr(v)})"

    def o(v):
        return f"Obs!{L(m.CAMPOS.index(v) + 1)}{rs}"

    linhas = []
    for v in m.ML_Y:
        linhas.append(("(3.14)", v, f"={sp(v)}", f"={o(v)}+{beta}*{o(v)}", "≥"))
    for v in m.ML_X:
        linhas.append(("(3.14)", v, f"={sp(v)}", f"={o(v)}-{beta}*{o(v)}", "≤"))
    for v in m.ML_B:
        linhas.append(("(3.14)", v, f"={sp(v)}", f"={o(v)}-{beta}*{o(v)}", "="))
    return linhas


def _restricoes(ws, regs):
    _cab(ws, 1, ["modelo", "sigla", "t", "a", "b", "eq.", "variável", "sentido",
                 "lado esquerdo", "lado direito", "folga"])
    k = 2
    for reg in regs:
        linhas = _linhas_ml(reg) if reg["kind"] == "ml" else _linhas_rede(reg)
        for eq, var, lhs, rhs, sentido in linhas:
            ws.cell(k, 1, reg["modelo"])
            ws.cell(k, 2, reg["cod"])
            ws.cell(k, 3, reg["t"])
            ws.cell(k, 4, reg["a"])
            ws.cell(k, 5, reg["b"])
            ws.cell(k, 6, eq)
            ws.cell(k, 7, var)
            ws.cell(k, 8, sentido)
            ws.cell(k, 9, lhs)
            ws.cell(k, 10, rhs)
            ws.cell(k, 11, f"=J{k}-I{k}" if sentido == "≤" else f"=I{k}-J{k}")
            k += 1


def _condicao16(ws, prep):
    _cab(ws, 1, ["par", "vetor", "variável", "exigido", "valor em t", "valor em t+1", "atende"])
    sinal = {"mu": "<", "omega": "<", "alpha": "<", "rho": ">", "nu": ">", "gamma": ">"}
    anos = prep["anos"]
    r = 2
    for it in range(len(anos) - 1):
        t = int(anos[it])
        for grupo, var in m.ORDEM_DIR:
            ws.cell(r, 1, f"{t}-{int(anos[it + 1])}")
            ws.cell(r, 2, grupo)
            ws.cell(r, 3, var)
            ws.cell(r, 4, sinal[grupo])
            for col, a in ((5, 0), (6, 1)):
                ws.cell(r, col,
                        f'=SUMIFS(Direcoes!$H:$H,Direcoes!$A:$A,"DSMLPI",Direcoes!$B:$B,{t},'
                        f'Direcoes!$C:$C,{a},Direcoes!$E:$E,C{r})')
            ws.cell(r, 7, f'=IF(D{r}="<",F{r}<E{r},F{r}>E{r})')
            r += 1


def _indices(ws, prep, lookup):
    _cab(ws, 1, ["modelo", "sigla", "ano (t+1)", "1+D^t(t)", "1+D^t(t+1)",
                 "1+D^{t+1}(t)", "1+D^{t+1}(t+1)", "EC", "TC", "índice", "ln EC", "ln TC", "ln índice"])
    anos = prep["anos"]
    i = 2
    for modelo in ("MLPI", "SMLPI", "DMLPI", "DSMLPI"):
        for cod in _bancos_do_modelo(prep, modelo):
            for it in range(len(anos) - 1):
                t = int(anos[it])
                ws.cell(i, 1, modelo)
                ws.cell(i, 2, cod)
                ws.cell(i, 3, int(anos[it + 1]))
                for col, (a, b) in zip((4, 5, 6, 7), ((0, 0), (0, 1), (1, 0), (1, 1))):
                    ld = lookup[(modelo, cod, t, a, b)]
                    ws.cell(i, col, f'=IF(Solucao!G{ld}<>0,"",1+Solucao!F{ld})')
                ok = f"AND(COUNT(D{i}:G{i})=4,MIN(D{i}:G{i})>0)"
                ws.cell(i, 8, f'=IF({ok},D{i}/G{i},"")')
                ws.cell(i, 9, f'=IF({ok},SQRT((G{i}/E{i})*(F{i}/D{i})),"")')
                ws.cell(i, 10, f'=IF({ok},SQRT((D{i}/E{i})*(F{i}/G{i})),"")')
                for col, src in ((11, "H"), (12, "I"), (13, "J")):
                    ws.cell(i, col, f'=IF({src}{i}="","",LN({src}{i}))')
                i += 1
    return i - 1


def _tabelas(ws, prep, n, norm, col_norm, n_painel):
    R = lambda c: f"Indices!${c}$2:${c}${n}"
    lin = 1
    ws.cell(lin, 1, "Tabela 2 · estatísticas descritivas").font = NEG
    _cab(ws, lin + 1, ["Variável", "Média", "σ", "Curtose", "Assimetria", "Mediana", "Máximo", "Mínimo"])
    r = lin + 2
    for v in pb.VARS_MODELO:
        rng = f"Painel!{col_norm[v]}$2:{col_norm[v]}${n_painel}"
        ws.cell(r, 1, pb.VARIAVEIS[v][0])
        for c, fn in enumerate(("AVERAGE", "_xlfn.STDEV.S", "KURT", "SKEW", "MEDIAN", "MAX", "MIN"), 2):
            ws.cell(r, c, f"={fn}({rng})")
        r += 1
    lin = r + 1
    blocos = [
        ("MLPI, SMLPI, DMLPI e DSMLPI por ano", ["MLPI", "SMLPI", "DMLPI", "DSMLPI"]),
    ]
    anos = [int(a) for a in prep["anos"][1:]]
    for titulo, modelos in blocos:
        ws.cell(lin, 1, titulo).font = NEG
        cab = ["Ano"]
        for mo in modelos:
            cab += [f"{mo} EC", f"{mo} TC", mo, f"{mo} sem solução"]
        _cab(ws, lin + 1, cab)
        rr = lin + 2
        for rot in [str(a) for a in anos] + ["período"]:
            ws.cell(rr, 1, rot)
            c = 2
            for mo in modelos:
                crit = f'{R("A")},"{mo}"'
                if rot != "período":
                    crit += f',{R("C")},{rot}'
                for lncol in ("K", "L", "M"):
                    ws.cell(rr, c, f'=IFERROR(EXP(AVERAGEIFS({R(lncol)},{crit})),"")')
                    c += 1
                ws.cell(rr, c, f'=COUNTIFS({crit},{R("J")},"")')
                c += 1
            rr += 1
        lin = rr + 1
    ws.cell(lin, 1, "Média geométrica por banco no período").font = NEG
    cab = ["sigla"]
    for mo in ("MLPI", "SMLPI", "DMLPI", "DSMLPI"):
        cab += [f"{mo} EC", f"{mo} TC", mo, f"{mo} sem solução"]
    _cab(ws, lin + 1, cab)
    rr = lin + 2
    for sigla in sorted(norm["sigla"].drop_duplicates()):
        ws.cell(rr, 1, sigla)
        c = 2
        for mo in ("MLPI", "SMLPI", "DMLPI", "DSMLPI"):
            crit = f'{R("A")},"{mo}",{R("B")},"{sigla}"'
            for lncol in ("K", "L", "M"):
                ws.cell(rr, c, f'=IFERROR(EXP(AVERAGEIFS({R(lncol)},{crit})),"")')
                c += 1
            ws.cell(rr, c, f'=COUNTIFS({crit},{R("J")},"")')
            c += 1
        rr += 1
    ws.column_dimensions["A"].width = 42


def gerar_excel(caminho, contas_long, norm, defl, ag):
    """Planilha com fórmulas. A aba Solucao liga a solucoes.xlsx."""
    norm = norm.sort_values(["ano", "sigla"]).reset_index(drop=True)
    prep = m.preparar(norm)
    wb = Workbook()
    _leiame(wb.active)
    ordem = list(norm[["ano", "sigla"]].itertuples(index=False, name=None))
    linha_de = {k: i + 2 for i, k in enumerate(ordem)}
    col_conta = _contas(wb.create_sheet("Contas"), contas_long, ordem)
    _dicionario(wb.create_sheet("Descricao"))
    _deflator(wb.create_sheet("Deflator"), defl)
    _agencias(wb.create_sheet("Agencias"), ag)
    col_norm = _painel(wb.create_sheet("Painel"), norm, ordem, col_conta)
    fr, dr = _fronteiras(wb, prep, linha_de, col_norm)
    lookup, regs = _solucao(wb, prep, fr, dr)
    _obs(wb, prep, lookup, linha_de, col_norm)
    _restricoes(wb.create_sheet("Restricoes"), regs)
    _condicao16(wb.create_sheet("Condicao16"), prep)
    n_ind = _indices(wb.create_sheet("Indices"), prep, lookup)
    _tabelas(wb.create_sheet("Tabelas"), prep, n_ind, norm, col_norm, len(ordem) + 1)
    for ws in wb.worksheets:
        ws.sheet_view.showGridLines = False
        ws.freeze_panes = "A2"
        for row in ws.iter_rows():
            for cel in row:
                al = cel.alignment
                if al is not None and al.wrap_text:
                    cel.alignment = Alignment(wrap_text=False, vertical=al.vertical or "bottom",
                                              horizontal=al.horizontal)
    wb.save(caminho)
    _anexar_ligacao(caminho)


_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS_REL = "http://schemas.openxmlformats.org/package/2006/relationships"
_NS_CT = "http://schemas.openxmlformats.org/package/2006/content-types"

def _xml(raiz):
    return ET.tostring(raiz, encoding="utf-8", xml_declaration=True)



def _numero(v):
    if v == "" or v is None:
        return None
    try:
        x = float(v)
    except (TypeError, ValueError):
        return None
    if not np.isfinite(x):
        return None
    return x


def _cache_fronteiras(prep):
    """Fronteira e direção por (modelo, t, a), na mesma ordem da aba Fronteiras."""
    anos = prep["anos"]
    cache = {}
    for modelo in ("DMLPI", "DSMLPI", "MLPI", "SMLPI"):
        seq = _modelo_seq(modelo)
        js = m.amostra_tabela5(prep)[0] if modelo in ("MLPI", "SMLPI") else None
        for it in range(len(anos) - 1):
            t = int(anos[it])
            for a in (0, 1):
                if js is not None:
                    idx = [it + a] if modelo == "MLPI" else list(range(it + a + 1))
                    fr = {v: np.concatenate([prep["dados"][v][i][js] for i in idx]) for v in m.CAMPOS}
                    cache[(modelo, t, a)] = fr, None
                else:
                    idx = list(range(it + a + 1)) if seq else [it + a]
                    fr = m.fronteira(prep, idx)
                    mult, sub = _constantes(modelo, a)
                    cache[(modelo, t, a)] = fr, m.direcoes(fr, mult, sub)
    return cache


def _matriz(fr):
    n = len(fr["labor"])
    return [[float(fr[v][i]) for v in m.CAMPOS] for i in range(n)]


def gravar_solucoes(caminho, prep):
    """Grava β, status e pesos como número. Não altera bansal.xlsx."""
    if os.path.basename(caminho).lower() == "bansal.xlsx":
        raise SystemExit("Este comando grava solucoes.xlsx. Não altera o bansal.xlsx.")
    cache = _cache_fronteiras(prep)
    anos = [int(a) for a in prep["anos"]]
    banco = {cod: j for j, cod in enumerate(prep["bancos"])}
    linhas = _linhas_lp(prep)
    wb = Workbook()
    ws = wb.active
    ws.title = "Solucao"
    _cab(ws, 1, ["modelo", "sigla", "t", "a", "b", "β", "status", "pesos"])
    total = len(linhas)
    for i, (modelo, cod, t, a, b) in enumerate(linhas, start=2):
        it = anos.index(t)
        fr, direcao = cache[(modelo, t, a)]
        kind = "ml" if modelo in ("MLPI", "SMLPI") else "rede"
        lista_d = ([0.0] * len(m.ORDEM_DIR) if direcao is None
                   else [direcao[g][v] for g, v in m.ORDEM_DIR])
        o = m.observacao(prep, it + b, banco[cod])
        res = m.otimo(kind, _matriz(fr), lista_d, [o[v] for v in m.CAMPOS])
        ws.cell(i, 1, modelo)
        ws.cell(i, 2, cod)
        ws.cell(i, 3, t)
        ws.cell(i, 4, a)
        ws.cell(i, 5, b)
        for c, v in enumerate(res, start=6):
            n = _numero(v)
            if n is not None:
                ws.cell(i, c, n)
        feito = i - 1
        if feito == 1 or feito % 25 == 0 or feito == total:
            print(f"{feito}/{total}  {modelo}  {t}", flush=True)
    try:
        wb.save(caminho)
    except PermissionError:
        raise SystemExit(f"{caminho} está aberto. Feche o arquivo e rode de novo.")


def _norm_local():
    con = pb.abrir_fonte()
    con_b = pb.abrir_bansal()
    try:
        _amo, _anuais, _nominal, norm, _ag, _defl, problemas, _zeros = pb.montar_painel(con, con_b)
    finally:
        con.close()
        con_b.close()
    if problemas or norm is None:
        texto = "\n".join(problemas or ["painel vazio"])
        raise SystemExit("Pendências na base:\n" + texto)
    return norm


def solucionar(caminho=None):
    if caminho is None:
        caminho = os.path.join(os.path.dirname(os.path.abspath(__file__)), ARQUIVO_SOLUCOES)
    norm = _norm_local()
    gravar_solucoes(caminho, m.preparar(norm))
    print(caminho, flush=True)


def main(argv):
    if len(argv) >= 2 and argv[1] == "--solucionar":
        solucionar(argv[2] if len(argv) > 2 else None)
        return 0
    print("Uso: python excel_bansal.py --solucionar [solucoes.xlsx]")
    return 2


def _anexar_ligacao(caminho):
    """Liga as fórmulas [1]Solucao! ao arquivo solucoes.xlsx, na mesma pasta."""
    with zipfile.ZipFile(caminho, "r") as zin:
        partes = {nome: zin.read(nome) for nome in zin.namelist()}
    rid, rels = _rels_ligacao(partes["xl/_rels/workbook.xml.rels"])
    partes["xl/_rels/workbook.xml.rels"] = rels
    partes["xl/workbook.xml"] = _workbook_ligacao(partes["xl/workbook.xml"], rid)
    partes["xl/externalLinks/externalLink1.xml"] = _external_link_xml()
    partes["xl/externalLinks/_rels/externalLink1.xml.rels"] = _external_rels()
    partes["[Content_Types].xml"] = _tipos_ligacao(partes["[Content_Types].xml"])
    tmp = str(caminho) + ".tmp"
    with zipfile.ZipFile(tmp, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for nome, dados in partes.items():
            zout.writestr(nome, dados)
    os.replace(tmp, caminho)


def _rels_ligacao(xml):
    ET.register_namespace("", _NS_REL)
    raiz = ET.fromstring(xml)
    ids = []
    for el in raiz:
        rid = el.get("Id") or ""
        if rid.startswith("rId") and rid[3:].isdigit():
            ids.append(int(rid[3:]))
    rid = f"rId{max(ids, default=0) + 1}"
    ET.SubElement(raiz, f"{{{_NS_REL}}}Relationship", {
        "Id": rid,
        "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLink",
        "Target": "externalLinks/externalLink1.xml",
    })
    return rid, _xml(raiz)


def _workbook_ligacao(xml, rid):
    texto = xml.decode("utf-8")
    if "externalReferences" not in texto:
        bloco = f'<externalReferences><externalReference r:id="{rid}"/></externalReferences>'
        if "</sheets>" not in texto:
            raise RuntimeError("workbook sem a lista de abas")
        texto = texto.replace("</sheets>", "</sheets>" + bloco, 1)
    return texto.encode("utf-8")


def _external_link_xml():
    ET.register_namespace("", _NS)
    ET.register_namespace("r", "http://schemas.openxmlformats.org/officeDocument/2006/relationships")
    raiz = ET.Element(f"{{{_NS}}}externalLink")
    book = ET.SubElement(raiz, f"{{{_NS}}}externalBook")
    book.set("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id", "rId1")
    nomes = ET.SubElement(book, f"{{{_NS}}}sheetNames")
    ET.SubElement(nomes, f"{{{_NS}}}sheetName", {"val": "Solucao"})
    dados = ET.SubElement(book, f"{{{_NS}}}sheetDataSet")
    ET.SubElement(dados, f"{{{_NS}}}sheetData", {"sheetId": "0"})
    return _xml(raiz)


def _external_rels():
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<Relationships xmlns="{_NS_REL}">'
        '<Relationship Id="rId1" '
        'Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/externalLinkPath" '
        f'Target="{ARQUIVO_SOLUCOES}" TargetMode="External"/>'
        "</Relationships>"
    ).encode("utf-8")


def _tipos_ligacao(xml):
    ET.register_namespace("", _NS_CT)
    raiz = ET.fromstring(xml)
    parte = "/xl/externalLinks/externalLink1.xml"
    if not any(el.get("PartName") == parte for el in raiz):
        ET.SubElement(raiz, f"{{{_NS_CT}}}Override", {
            "PartName": parte,
            "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.externalLink+xml",
        })
    return _xml(raiz)


if __name__ == "__main__":
    sys.exit(main(sys.argv))
